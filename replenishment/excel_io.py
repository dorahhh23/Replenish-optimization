"""Read-only Excel inspection, merged headers, semantic field identification."""
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
import pandas as pd
from openpyxl.utils.cell import range_boundaries
from .periods import parse_period_label
from .status import normalize_product_status


class SchemaError(ValueError):
    """Formal analysis must pause until the source mapping is unambiguous."""


ALIASES = {
    "material": {"物料", "物料编码", "sku", "material", "material code"},
    "sales_org": {"销售组织", "sales organization", "sales org"},
    "storage_location": {"库位", "storage location"},
    "product_status": {"产品状态", "product status", "product_status"},
    "product_category": {"国贸产品分类码", "product category"},
    "hierarchy_level": {"级别", "hierarchy level"},
    "unit": {"基本计量单位", "销售单位", "计量单位", "unit", "uom", "base unit"},
    "date": {"日期", "销售日期", "date", "sales date", "week start", "month"},
    "sales": {"销量", "销售数量", "sales", "quantity", "demand"},
}


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def identifier(value, number_format="General"):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)) or float(value) != int(value):
            return "<INVALID_NUMERIC_ID>"
        text = str(int(value))
        return text.zfill(len(number_format)) if re.fullmatch(r"0+", number_format) else text
    return str(value).strip()


def merged_ranges(path):
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with zipfile.ZipFile(path) as archive:
        rel = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {n.attrib["Id"]: n.attrib["Target"] for n in rel}
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        result = {}
        for node in book.find("m:sheets", ns):
            target = targets[node.attrib[f"{{{ns['r']}}}id"]]
            location = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            # Iterparse avoids building the entire large worksheet XML tree.
            ranges = []
            with archive.open(location) as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag.endswith("}mergeCell"):
                        ranges.append(element.attrib["ref"])
                    element.clear()
            result[node.attrib["name"]] = ranges
        return result


def parse_headers(raw_headers, merges=()):
    grid = [list(row) for row in raw_headers]
    for address in merges:
        c1, r1, c2, r2 = range_boundaries(address)
        if r1 <= len(grid):
            value = grid[r1 - 1][c1 - 1]
            for row in range(r1 - 1, min(r2, len(grid))):
                for col in range(c1 - 1, min(c2, len(grid[row]))):
                    if grid[row][col] is None:
                        grid[row][col] = value
    used = Counter()
    fields = []
    for col in range(len(grid[0])):
        parts = []
        raw = []
        for row in grid:
            value = row[col]
            if value is None or str(value).strip() == "":
                continue
            raw.append(value)
            label = value.isoformat() if isinstance(value, (date, datetime)) else re.sub(r"\s+", " ", str(value).strip())
            if not parts or label != parts[-1]:
                parts.append(label)
        base = " | ".join(parts) or f"Unnamed column {col + 1}"
        used[base] += 1
        name = base if used[base] == 1 else f"{base} [{used[base]}]"
        fields.append({"column_index": col + 1, "original_name": " | ".join(map(str, raw)),
                       "cleaned_name": name, "group": parts[0] if len(parts) > 1 else "",
                       "leaf": parts[-1] if parts else "", "raw_leaf": raw[-1] if raw else None})
    return fields


def read_workbook(path, config):
    path = Path(path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Source workbook not found: {path}")
    digest = sha256(path)
    merges = merged_ranges(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables, dictionary, structure = {}, [], []
    try:
        for sheet in book:
            iterator = sheet.iter_rows()
            headers = [tuple(c.value for c in next(iterator)) for _ in range(config.header_rows)]
            fields = parse_headers(headers, merges.get(sheet.title, []))
            id_cols = {f["column_index"] - 1 for f in fields
                       if f["leaf"].casefold() in set().union(*[ALIASES[k] for k in ("material", "sales_org", "storage_location")])}
            values, source_rows, numeric_id_cells = [], [], 0
            for source_row, cells in enumerate(iterator, config.header_rows + 1):
                row = [cell.value for cell in cells]
                if not any(value is not None for value in row):
                    continue
                for col in id_cols:
                    cell = cells[col]
                    numeric_id_cells += isinstance(cell.value, (int, float))
                    row[col] = identifier(cell.value, cell.number_format)
                values.append(row)
                source_rows.append(source_row)
            data = pd.DataFrame(values, columns=[f["cleaned_name"] for f in fields])
            data["_source_row"] = source_rows
            tables[sheet.title] = (data, fields)
            structure.append({"sheet": sheet.title, "excel_rows": sheet.max_row,
                              "data_rows": len(data), "columns": len(fields),
                              "header_rows": config.header_rows, "merged_headers": merges.get(sheet.title, []),
                              "numeric_identifier_cells": numeric_id_cells})
            for f in fields:
                s = data[f["cleaned_name"]]
                missing = s.isna() | s.astype(str).str.strip().eq("")
                examples = s.loc[~missing].drop_duplicates().head(4).tolist()
                dictionary.append({**{k: v for k, v in f.items() if k != "raw_leaf"}, "sheet": sheet.title,
                                   "data_type": str(s.dtype), "source_value_types": ", ".join(sorted({type(x).__name__ for x in s.dropna()})),
                                   "missing_rate": float(missing.mean()) if len(s) else None,
                                   "example_values": json.dumps(examples, ensure_ascii=False, default=str),
                                   "meaning": f["leaf"], "used_for_analysis": False, "purpose": "inspection only"})
    finally:
        book.close()
    if sha256(path) != digest:
        raise RuntimeError("Source workbook changed while it was being read.")
    return tables, pd.DataFrame(dictionary), structure, digest


def identify_fields(data, fields, config):
    mapping, diagnostics = {}, []
    for role, aliases in ALIASES.items():
        candidates = []
        for f in fields:
            s = data[f["cleaned_name"]]
            exact = f["leaf"].casefold() in aliases
            status_like = role == "product_status" and bool(re.search(r"状态|status", f["leaf"], re.I))
            if exact or status_like:
                unique = s.dropna().drop_duplicates()
                normalized = unique.map(normalize_product_status)
                permitted = int(normalized.isin(["Normal", "Phasing Out"]).sum())
                diagnostics.append({"role": role, "field": f["cleaned_name"], "exact_semantic_match": exact,
                                    "unique_count": len(unique), "unique_values": json.dumps(unique.head(100 if status_like else 20).tolist(), ensure_ascii=False, default=str),
                                    "unique_values_truncated": len(unique) > (100 if status_like else 20),
                                    "examples": json.dumps(unique.head(5).tolist(), ensure_ascii=False, default=str),
                                    "allowed_status_values": permitted if status_like else None})
                if exact or (status_like and permitted > 0):
                    candidates.append(f["cleaned_name"])
        if role in config.fields:
            selected = config.fields[role]
            if selected not in data:
                raise SchemaError(f"Configured {role} field is absent: {selected}")
            mapping[role] = selected
        elif len(candidates) == 1:
            mapping[role] = candidates[0]
        elif len(candidates) > 1 and role in {"material", "sales_org", "storage_location", "product_status", "date", "sales", "unit"}:
            raise SchemaError(f"Multiple possible fields for {role}: {candidates}. Set fields explicitly after inspection.")
    for required in ("material", "product_status"):
        if required not in mapping:
            raise SchemaError(f"Missing required field: {required}. Review field_candidates.json.")
    return mapping, pd.DataFrame(diagnostics)


def field_candidates(data, fields):
    """Always write status candidates before strict identification can raise."""
    return [{"field": f["cleaned_name"], "unique_values": data[f["cleaned_name"]].dropna().drop_duplicates().head(100).tolist(),
             "examples": data[f["cleaned_name"]].dropna().head(5).tolist()}
            for f in fields if re.search(r"状态|status", f["leaf"], re.I)]


def identify_period_columns(fields, actual_group):
    periods, unknown = [], []
    for f in fields:
        if f["group"].casefold() != actual_group.casefold():
            continue
        try:
            parsed = parse_period_label(f["raw_leaf"])
        except ValueError as error:
            unknown.append({"field": f["cleaned_name"], "reason": str(error)})
            continue
        if parsed:
            periods.append({"field": f["cleaned_name"], **parsed})
        else:
            unknown.append({"field": f["cleaned_name"], "reason": "No unambiguous year/date. No year inferred from column order."})
    return periods, unknown
